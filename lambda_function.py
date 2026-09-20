import boto3
import urllib.parse
from openpyxl import load_workbook

s3 = boto3.client("s3")
dynamodb = boto3.resource("dynamodb")

calendar_table = dynamodb.Table("Calendar")
users_table = dynamodb.Table("Users")


def lambda_handler(event, context):
    print("EVENT:", event)

    record = event["Records"][0]

    bucket = record["s3"]["bucket"]["name"]
    key = urllib.parse.unquote_plus(
        record["s3"]["object"]["key"]
    )

    print("bucket:", bucket)
    print("key:", key)

    local_path = "/tmp/upload.xlsx"

    # S3에서 Excel 다운로드
    s3.download_file(
        bucket,
        key,
        local_path
    )

    workbook = load_workbook(
        local_path,
        data_only=True
    )

    print("sheets:", workbook.sheetnames)

    # 지금은 시트 하나만 남긴 상태라고 가정
    sheet = workbook[workbook.sheetnames[0]]

    header_row = None
    columns = {}

    # 위에서부터 10행 안에서 헤더 찾기
    for row_index in range(1, 11):

        values = [
            sheet.cell(row=row_index, column=col).value
            for col in range(1, sheet.max_column + 1)
        ]

        if "日付" in values and "担当者" in values:
            header_row = row_index

            for col_index, value in enumerate(values, start=1):
                if value is not None:
                    columns[str(value)] = col_index

            break

    if header_row is None:
        raise Exception("日付 / 担当者 ヘッダーが見つかりません")

    print("header_row:", header_row)
    print("columns:", columns)

    # Users 전체 조회
    users_response = users_table.scan()
    users = users_response.get("Items", [])

    # 이름 -> userId
    user_map = {}

    for user in users:
        name = user.get("name")
        user_id = user.get("userId")

        if name and user_id:
            user_map[name] = user_id

    print("user_map:", user_map)

    saved_count = 0

    for row_index in range(
        header_row + 1,
        sheet.max_row + 1
    ):
        date_value = sheet.cell(
            row=row_index,
            column=columns["日付"]
        ).value

        if date_value is None:
            continue

        date_string = date_value.strftime("%Y-%m-%d")
        month_string = date_value.strftime("%Y-%m")

        assigned_name = None
        assigned_user_id = None

        if "担当者" in columns:
            assigned_name = sheet.cell(
                row=row_index,
                column=columns["担当者"]
            ).value

        if assigned_name:
            assigned_user_id = user_map.get(assigned_name)

        jwd = False

        if "JWD" in columns:
            value = sheet.cell(
                row=row_index,
                column=columns["JWD"]
            ).value

            jwd = value == "〇"

        holiday_work_request = False

        if "休出申請" in columns:
            value = sheet.cell(
                row=row_index,
                column=columns["休出申請"]
            ).value

            holiday_work_request = value == "〇"

        holiday_name = None

        if "備考" in columns:
            holiday_name = sheet.cell(
                row=row_index,
                column=columns["備考"]
            ).value

        item = {
            "month": month_string,
            "date": date_string,
            "jwd": jwd,
            "holidayWorkRequest": holiday_work_request,
            "holiday": holiday_name is not None
        }

        if assigned_name:
            item["originalUserName"] = assigned_name
            item["assignedUserName"] = assigned_name

        if assigned_user_id:
            item["originalUserId"] = assigned_user_id
            item["assignedUserId"] = assigned_user_id

        if holiday_name:
            item["holidayName"] = holiday_name

        print("saving:", item)

        calendar_table.put_item(
            Item=item
        )

        saved_count += 1

    print("saved_count:", saved_count)

    return {
        "statusCode": 200,
        "savedCount": saved_count
    }